"""Create an editable Word report from the local authoritative Markdown source."""

import re
from pathlib import Path


def build(root):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from PIL import Image

    source = root / "reports/technical-report.md"
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "Correlated Detection of MITRE ATT&CK Sub-Techniques"
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(9.5)
        style.paragraph_format.line_spacing = Pt(12.2)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.space_before = Pt(0)
    for name, size in [("Title", 21), ("Heading 1", 13.5), ("Heading 2", 14)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.line_spacing = Pt(size + 3)
        style.paragraph_format.keep_with_next = True

    def inline(paragraph, text, size=None):
        for token in re.split(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                relation = paragraph.part.relate_to(
                    link[2],
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                    is_external=True,
                )
                element = OxmlElement("w:hyperlink")
                element.set(qn("r:id"), relation)
                run = OxmlElement("w:r")
                properties = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "195878")
                properties.append(color)
                run.append(properties)
                value = OxmlElement("w:t")
                value.text = link[1]
                run.append(value)
                element.append(run)
                paragraph._p.append(element)
                continue
            value = token[2:-2] if token.startswith("**") else token.strip("`")
            run = paragraph.add_run(value)
            if token.startswith("**"):
                run.bold = True
            elif token.startswith("*") and token.endswith("*"):
                run.text = token[1:-1]
                run.italic = True
            if token.startswith("`"):
                run.font.name = "Courier New"
                run.font.size = Pt(size or 8.5)
            elif size:
                run.font.size = Pt(size)

    parts = re.split(r"<!-- page: (portrait|landscape) -->", source.read_text(encoding="utf-8"))
    pages = list(zip(parts[1::2], parts[2::2], strict=True))
    for number, (orientation, markdown) in enumerate(pages, 1):
        section = doc.sections[0] if number == 1 else doc.add_section(WD_SECTION.NEW_PAGE)
        wide = orientation == "landscape"
        section.page_width = Pt(841.89 if wide else 595.28)
        section.page_height = Pt(595.28 if wide else 841.89)
        section.left_margin = section.right_margin = Pt(36)
        section.top_margin = Pt(28)
        section.bottom_margin = Pt(41)
        section.footer_distance = Pt(15)
        section.footer.is_linked_to_previous = False
        footer = section.footer.paragraphs[0]
        footer.paragraph_format.line_spacing = Pt(9)
        run = footer.add_run(
            f"ATT&CK detection and correlation | Windows-APT 2025 v4        {number} / {len(pages)}"
        )
        run.font.name = "Arial"
        run.font.size = Pt(8)
        width = section.page_width.pt - 72
        lines = markdown.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue
            if line.startswith("```"):
                i += 1
                code = []
                while i < len(lines) and not lines[i].startswith("```"):
                    code.append(lines[i])
                    i += 1
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.line_spacing = Pt(9.5)
                run = paragraph.add_run("\n".join(code))
                run.font.name = "Courier New"
                run.font.size = Pt(7.5)
            elif line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    values = re.split(r"(?<!\\)\|", lines[i].strip().strip("|"))
                    if not all(re.fullmatch(r"[\s:\-]+", v) for v in values):
                        rows.append([v.strip() for v in values])
                    i += 1
                weights = {
                    2: [0.35, 0.65],
                    3: [0.70, 0.15, 0.15],
                    5: [0.14, 0.20, 0.28, 0.18, 0.20],
                }[len(rows[0])]
                table = doc.add_table(rows=0, cols=len(weights))
                table.autofit = False
                for column, weight in zip(table.columns, weights, strict=True):
                    column.width = Pt(width * weight)
                borders = OxmlElement("w:tblBorders")
                for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                    border = OxmlElement("w:" + edge)
                    for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                for row_index, values in enumerate(rows):
                    row = table.add_row()
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    if row_index == 0:
                        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                    for cell, value, weight in zip(row.cells, values, weights, strict=True):
                        cell.width = Pt(width * weight)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        properties = cell._tc.get_or_add_tcPr()
                        margins = OxmlElement("w:tcMar")
                        for edge in ["top", "left", "bottom", "right"]:
                            margin = OxmlElement("w:" + edge)
                            margin.set(qn("w:w"), "90")
                            margin.set(qn("w:type"), "dxa")
                            margins.append(margin)
                        properties.append(margins)
                        if row_index == 0:
                            fill = OxmlElement("w:shd")
                            fill.set(qn("w:fill"), "DEEDF2")
                            properties.append(fill)
                        paragraph = cell.paragraphs[0]
                        paragraph.paragraph_format.line_spacing = Pt(10.5)
                        paragraph.paragraph_format.space_after = Pt(0)
                        inline(paragraph, value, 8.2)
                gap = doc.add_paragraph()
                gap.paragraph_format.space_after = Pt(0)
                gap.paragraph_format.line_spacing = Pt(4)
                continue
            elif line.startswith("!["):
                image_path = root / ".local/final/pdf-pages" / f"figure-{number}.png"
                with Image.open(image_path) as figure:
                    scale = min(width / figure.width, 418 / figure.height)
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.add_run().add_picture(str(image_path), width=Pt(figure.width * scale))
                    paragraph.paragraph_format.keep_with_next = True
                    drawing = paragraph._p.xpath(".//wp:docPr")[0]
                    drawing.set("descr", re.search(r"!\[([^\]]+)", line)[1])
            elif line.startswith("#"):
                match = re.match(r"(#{1,3})\s+(.*)", line)
                style = {1: "Title", 2: "Heading 1", 3: "Heading 2"}[len(match[1])]
                paragraph = doc.add_paragraph(style=style)
                inline(paragraph, match[2])
            else:
                paragraph_lines = [line]
                i += 1
                while (
                    i < len(lines)
                    and lines[i].strip()
                    and not lines[i].startswith(("#", "|", "```", "![", "- "))
                ):
                    paragraph_lines.append(lines[i].strip())
                    i += 1
                content = " ".join(paragraph_lines)
                paragraph = doc.add_paragraph()
                inline(paragraph, "• " + content[2:] if content.startswith("- ") else content)
                continue
            i += 1
    output = root / "reports/technical-report.docx"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build(Path(__file__).resolve().parents[1])
